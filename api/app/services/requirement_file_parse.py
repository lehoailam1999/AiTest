"""Extract plain text and optional HTML preview from requirement upload files."""

from __future__ import annotations

import base64
import html as html_module
import io
import re
import zipfile
from dataclasses import dataclass
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser


@dataclass(frozen=True)
class ParsedRequirementFile:
    text: str
    parser: str
    warning: str
    html: str | None = None


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._chunks: list[str] = []
        self._skip = False

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in ("script", "style"):
            self._skip = True
        elif tag in ("p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "td"):
            self._chunks.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("script", "style"):
            self._skip = False

    def handle_data(self, data: str) -> None:
        if not self._skip and data:
            self._chunks.append(data)

    def text(self) -> str:
        raw = "".join(self._chunks)
        raw = re.sub(r"[ \t]+\n", "\n", raw)
        raw = re.sub(r"\n{3,}", "\n\n", raw)
        return raw.strip()


def html_to_text(html: str) -> str:
    parser = _HTMLTextExtractor()
    try:
        parser.feed(html)
        parser.close()
    except Exception:
        return re.sub(r"<[^>]+>", " ", html)
    return parser.text()


def _wrap_preview_html(body: str) -> str:
    return f'<div class="req-rich-preview">{body}</div>'


def _decode_raw(raw: bytes) -> str | None:
    for enc in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return None


def _is_word_ole(raw: bytes) -> bool:
    return len(raw) >= 8 and raw[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _is_zip_office(raw: bytes) -> bool:
    return len(raw) >= 2 and raw[:2] == b"PK"


def extract_mime_multipart(raw: bytes) -> str | None:
    doc = extract_mime_document(raw)
    return doc.text if doc else None


def extract_mime_document(raw: bytes) -> ParsedRequirementFile | None:
    """Confluence / Word HTML MIME — text + HTML with embedded images."""
    try:
        msg = BytesParser(policy=policy.default).parsebytes(raw)
    except Exception:
        return None

    cid_urls: dict[str, str] = {}
    location_urls: dict[str, str] = {}
    html_chunks: list[str] = []
    plain_chunks: list[str] = []

    for part in msg.walk():
        if part.get_content_maintype() == "multipart":
            continue
        ctype = part.get_content_type()
        if ctype.startswith("image/"):
            payload = part.get_payload(decode=True)
            if not payload:
                continue
            url = f"data:{ctype};base64,{base64.b64encode(payload).decode('ascii')}"
            cid = part.get("Content-ID")
            if cid:
                key = cid.strip().strip("<>")
                cid_urls[key] = url
            loc = part.get("Content-Location")
            if loc:
                location_urls[loc.strip()] = url
        elif ctype == "text/html":
            try:
                payload = part.get_content()
                if isinstance(payload, str) and payload.strip():
                    html_chunks.append(payload)
            except Exception:
                continue
        elif ctype == "text/plain":
            try:
                payload = part.get_content()
                if isinstance(payload, str) and payload.strip():
                    plain_chunks.append(payload.strip())
            except Exception:
                continue

    html = max(html_chunks, key=len) if html_chunks else ""
    if html:
        for cid, url in cid_urls.items():
            html = html.replace(f"cid:{cid}", url)
            html = re.sub(
                rf'src=(["\'])cid:{re.escape(cid)}\1',
                f'src="{url}"',
                html,
                flags=re.I,
            )
        for loc, url in location_urls.items():
            html = html.replace(loc, url)

        text = html_to_text(html)
        if not text.strip() and plain_chunks:
            text = "\n\n".join(plain_chunks)
        if not text.strip():
            return None
        return ParsedRequirementFile(
            text=text.strip(),
            parser="mime-doc",
            warning="",
            html=_wrap_preview_html(html),
        )

    if plain_chunks:
        text = "\n\n".join(plain_chunks)
        safe = html_module.escape(text)
        return ParsedRequirementFile(
            text=text,
            parser="mime-plain",
            warning="",
            html=_wrap_preview_html(f"<pre>{safe}</pre>"),
        )
    return None


def parse_docx_bytes(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts: list[str] = []
    for para in doc.paragraphs:
        t = (para.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts).strip()


def parse_docx_document(raw: bytes) -> ParsedRequirementFile:
    text = parse_docx_bytes(raw)
    if not text.strip():
        raise ValueError("File Word không có nội dung text — thử export .txt hoặc paste thủ công.")

    body_parts: list[str] = []
    for block in text.split("\n\n"):
        block = block.strip()
        if block:
            body_parts.append(f"<p>{html_module.escape(block)}</p>")

    imgs: list[str] = []
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            for name in sorted(zf.namelist()):
                if not name.startswith("word/media/"):
                    continue
                ext = name.rsplit(".", 1)[-1].lower()
                if ext not in ("png", "jpg", "jpeg", "gif", "webp"):
                    continue
                data = zf.read(name)
                mime = {
                    "png": "image/png",
                    "jpg": "image/jpeg",
                    "jpeg": "image/jpeg",
                    "gif": "image/gif",
                    "webp": "image/webp",
                }[ext]
                url = f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}"
                alt = html_module.escape(name.split("/")[-1])
                imgs.append(f'<figure><img src="{url}" alt="{alt}"/></figure>')
    except Exception:
        pass

    if imgs:
        body_parts.append('<section class="doc-images">' + "".join(imgs) + "</section>")

    html = _wrap_preview_html("".join(body_parts))
    return ParsedRequirementFile(text=text, parser="docx", warning="", html=html)


def _looks_like_mime_envelope(text: str) -> bool:
    head = text[:4000]
    return "MIME-Version:" in head and "Content-Type:" in head


def _mostly_headers(text: str) -> bool:
    lines = [ln for ln in text.splitlines()[:40] if ln.strip()]
    if len(lines) < 3:
        return False
    headerish = sum(1 for ln in lines if re.match(r"^[A-Za-z0-9-]+:\s", ln))
    return headerish >= max(3, len(lines) // 2)


def parse_requirement_file(raw: bytes, filename: str) -> ParsedRequirementFile:
    name = filename or "upload"
    ext = ("." + name.rsplit(".", 1)[-1].lower()) if "." in name else ""
    warning = ""

    text_ext = {
        ".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".xml",
        ".html", ".htm", ".cs", ".ts", ".js", ".go", ".py", ".sql", ".feature",
    }

    if ext == ".docx" or (ext == ".doc" and _is_zip_office(raw)):
        try:
            return parse_docx_document(raw)
        except ValueError:
            raise
        except Exception as e:
            raise ValueError(f"Không đọc được file Word (.docx): {e}") from e

    if ext == ".doc":
        if _is_word_ole(raw):
            raise ValueError(
                "File .doc dạng Word cũ (binary). Hãy mở bằng Word/LibreOffice và "
                "Lưu thành .docx, hoặc export PDF / copy text vào ô Yêu cầu."
            )
        mime_doc = extract_mime_document(raw)
        if mime_doc:
            return mime_doc
        decoded = _decode_raw(raw)
        if decoded and _looks_like_mime_envelope(decoded):
            mime_doc = extract_mime_document(decoded.encode("latin-1", errors="ignore"))
            if mime_doc:
                return mime_doc
        if decoded and not _mostly_headers(decoded):
            safe = html_module.escape(decoded.strip())
            return ParsedRequirementFile(
                text=decoded.strip(),
                parser="doc-text",
                warning="File .doc dạng text — đã đọc best-effort; nên dùng .docx nếu có thể.",
                html=_wrap_preview_html(f"<pre>{safe}</pre>"),
            )
        raise ValueError(
            "Không trích xuất được nội dung từ .doc. "
            "Export từ Confluence/Word sang .docx hoặc .md/.txt, hoặc paste text."
        )

    if ext in text_ext:
        decoded = _decode_raw(raw)
        if decoded is None:
            raise ValueError("File không phải UTF-8 — lưu lại UTF-8 hoặc dùng .docx.")
        if ext in (".html", ".htm"):
            text = html_to_text(decoded) or decoded.strip()
            return ParsedRequirementFile(
                text=text,
                parser="html",
                warning=warning,
                html=_wrap_preview_html(decoded),
            )
        safe = html_module.escape(decoded.strip())
        return ParsedRequirementFile(
            text=decoded.strip(),
            parser="text",
            warning=warning,
            html=_wrap_preview_html(f"<pre>{safe}</pre>"),
        )

    if ext in {".pdf", ".xlsx", ".xls", ".zip"}:
        raise ValueError(
            f"Chưa hỗ trợ parse {ext} — dùng .docx, .md, .txt hoặc paste nội dung."
        )

    decoded = _decode_raw(raw)
    if decoded and decoded.strip():
        if _looks_like_mime_envelope(decoded):
            mime_doc = extract_mime_document(raw if raw else decoded.encode("utf-8", errors="ignore"))
            if mime_doc:
                return mime_doc
        if _mostly_headers(decoded):
            raise ValueError("File giống email/MIME — export lại dạng .docx hoặc .txt.")
        safe = html_module.escape(decoded.strip())
        return ParsedRequirementFile(
            text=decoded.strip(),
            parser="utf8-fallback",
            warning=f"Đuôi file lạ ({ext or 'không có'}) — đọc như text UTF-8.",
            html=_wrap_preview_html(f"<pre>{safe}</pre>"),
        )

    raise ValueError(f"Không đọc được nội dung từ «{name}».")
